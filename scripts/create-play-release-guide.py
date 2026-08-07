from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Iterator

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
RELEASE_DIR = ROOT / "docs" / "release"

PRESET_NAME = "compact_reference_guide"
HEADER_TEMPLATE = "editorial_cover"
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
TABLE_ROW_KEEP_TOGETHER_MAX_CHARS = 180

BLUE = "2E74B5"
NAVY = "203748"
DEEP_BLUE = "1F4D78"
PALE_BLUE = "E8EEF5"
PALE_ROSE = "FFF0F5"
ROSE = "B0477A"
GOLD = "B7862D"
MUTED = "667085"
GRID = "C9D3DF"


@dataclass(frozen=True)
class GuideMetadata:
    version: str
    version_code: str
    written_date: str
    package_name: str


def find_default_source(release_dir: Path = RELEASE_DIR) -> Path:
    candidates = sorted(release_dir.glob("혜니캘린더_Google_Play_출시_가이드북_*.md"))
    if not candidates:
        raise FileNotFoundError(f"출시 가이드 Markdown을 찾을 수 없습니다: {release_dir}")
    return candidates[-1]


def parse_metadata(text: str) -> GuideMetadata:
    version_match = re.search(r"버전\s+([^\s·]+)\s*·\s*versionCode\s+(\d+)", text)
    date_match = re.search(r"작성일\s+(\d{4}-\d{2}-\d{2})", text)
    package_match = re.search(r"패키지\s+`([^`]+)`", text)
    if not version_match or not date_match or not package_match:
        raise ValueError("가이드 상단의 버전, versionCode, 작성일, 패키지 메타데이터를 확인해 주세요.")
    return GuideMetadata(
        version=version_match.group(1),
        version_code=version_match.group(2),
        written_date=date_match.group(1),
        package_name=package_match.group(1),
    )


def is_cover_metadata_line(stripped: str) -> bool:
    return (
        stripped.startswith("# ")
        or stripped.startswith("버전 ")
        or stripped.startswith("작성일 ")
    )


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_run_font(run, name: str = "Calibri", size: float | None = None) -> None:
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:eastAsia"), "Malgun Gothic")
    if size is not None:
        run.font.size = Pt(size)


def shade(element, fill: str) -> None:
    if hasattr(element, "_tc"):
        element = element._tc
    props = element.get_or_add_tcPr() if element.tag.endswith("tc") else element.get_or_add_pPr()
    shd = props.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        props.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(
    cell,
    top: int = 80,
    start: int = 120,
    bottom: int = 80,
    end: int = 120,
) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def replace_child(parent, tag: str) -> OxmlElement:
    existing = parent.find(qn(tag))
    if existing is not None:
        parent.remove(existing)
    node = OxmlElement(tag)
    parent.append(node)
    return node


def resolve_column_widths(column_count: int) -> list[int]:
    presets = {
        1: [9360],
        2: [2700, 6660],
        3: [2160, 3600, 3600],
        4: [1440, 2640, 2640, 2640],
        5: [1300, 1900, 1900, 2160, 2100],
    }
    if column_count in presets:
        return presets[column_count]
    base, remainder = divmod(CONTENT_WIDTH_DXA, column_count)
    return [base + (1 if index < remainder else 0) for index in range(column_count)]


def apply_table_geometry(table, widths_dxa: list[int]) -> None:
    if sum(widths_dxa) != CONTENT_WIDTH_DXA:
        raise ValueError(f"표 열 너비 합계가 {CONTENT_WIDTH_DXA} DXA가 아닙니다: {widths_dxa}")

    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = replace_child(tbl_pr, "w:tblW")
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = replace_child(tbl_pr, "w:tblInd")
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    tbl_layout = replace_child(tbl_pr, "w:tblLayout")
    tbl_layout.set(qn("w:type"), "fixed")

    tbl_grid = table._tbl.tblGrid
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        tbl_grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_dxa[min(index, len(widths_dxa) - 1)]
            tc_w = replace_child(cell._tc.get_or_add_tcPr(), "w:tcW")
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    set_run_font(run, size=9)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, end))


def add_hyperlink(paragraph, label: str, url: str, color: str = BLUE) -> None:
    rel = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_color = OxmlElement("w:color")
    r_color.set(qn("w:val"), color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Calibri")
    r_fonts.set(qn("w:hAnsi"), "Calibri")
    r_fonts.set(qn("w:eastAsia"), "Malgun Gothic")
    r_pr.extend((r_fonts, r_color, underline))
    run.append(r_pr)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


TOKEN = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|https?://[^\s)]+)")


def add_inline(paragraph, text: str) -> None:
    cursor = 0
    for match in TOKEN.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            set_run_font(run)
        token = match.group(0)
        if token.startswith("http"):
            add_hyperlink(paragraph, token, token)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, "Consolas", 9.5)
            run.font.color.rgb = rgb(DEEP_BLUE)
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), PALE_BLUE)
            run._element.get_or_add_rPr().append(shd)
        else:
            run = paragraph.add_run(token[2:-2])
            set_run_font(run)
            run.bold = True
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Title", 30, NAVY, 0, 8),
        ("Subtitle", 17, DEEP_BLUE, 0, 3),
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DEEP_BLUE, 10, 5),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
        style.font.size = Pt(size)
        style.font.color.rgb = rgb(color)
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    if "Code Block" not in styles:
        code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code = styles["Code Block"]
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "D2Coding")
    code.font.size = Pt(8.5)
    code.font.color.rgb = rgb(NAVY)
    code.paragraph_format.left_indent = Inches(0.12)
    code.paragraph_format.right_indent = Inches(0.12)
    code.paragraph_format.space_before = Pt(2)
    code.paragraph_format.space_after = Pt(4)
    code.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE


def configure_section(section, first: bool = False) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = first


def populate_header(header) -> None:
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run("혜니캘린더 · Google Play 출시 가이드북")
    set_run_font(run, size=8.5)
    run.font.color.rgb = rgb(MUTED)

    p_pr = paragraph._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), GRID)
    border.append(bottom)
    p_pr.append(border)


def populate_footer(footer) -> None:
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    label = footer_paragraph.add_run("페이지 ")
    set_run_font(label, size=9)
    label.font.color.rgb = rgb(MUTED)
    add_field(footer_paragraph, "PAGE")


def add_header_footer(section) -> None:
    populate_header(section.header)
    populate_footer(section.footer)
    populate_header(section.even_page_header)
    populate_footer(section.even_page_footer)


def configure_field_updates(doc: Document) -> None:
    settings = doc.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def create_abstract_numbering(numbering, abstract_id: int, kind: str) -> None:
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi_level = OxmlElement("w:multiLevelType")
    multi_level.set(qn("w:val"), "singleLevel")
    abstract.append(multi_level)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    level.append(num_fmt)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    level.append(level_text)
    suffix = OxmlElement("w:suff")
    suffix.set(qn("w:val"), "tab" if kind == "bullet" else "space")
    level.append(suffix)
    level_jc = OxmlElement("w:lvlJc")
    level_jc.set(qn("w:val"), "left")
    level.append(level_jc)

    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "540")
    indent.set(qn("w:hanging"), "270")
    p_pr.append(indent)
    level.append(p_pr)

    if kind == "bullet":
        r_pr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), "Calibri")
        fonts.set(qn("w:hAnsi"), "Calibri")
        r_pr.append(fonts)
        level.append(r_pr)

    abstract.append(level)
    numbering.append(abstract)


class NumberingManager:
    def __init__(self, doc: Document) -> None:
        self.numbering = doc.part.numbering_part.element
        abstract_ids = [
            int(node.get(qn("w:abstractNumId")))
            for node in self.numbering.findall(qn("w:abstractNum"))
        ]
        num_ids = [int(node.get(qn("w:numId"))) for node in self.numbering.findall(qn("w:num"))]
        next_abstract = max(abstract_ids, default=-1) + 1
        self.abstract_ids = {"bullet": next_abstract, "decimal": next_abstract + 1}
        create_abstract_numbering(self.numbering, self.abstract_ids["bullet"], "bullet")
        create_abstract_numbering(self.numbering, self.abstract_ids["decimal"], "decimal")
        self.next_num_id = max(num_ids, default=0) + 1

    def new_sequence(self, kind: str) -> int:
        num_id = self.next_num_id
        self.next_num_id += 1
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract = OxmlElement("w:abstractNumId")
        abstract.set(qn("w:val"), str(self.abstract_ids[kind]))
        num.append(abstract)
        level_override = OxmlElement("w:lvlOverride")
        level_override.set(qn("w:ilvl"), "0")
        start_override = OxmlElement("w:startOverride")
        start_override.set(qn("w:val"), "1")
        level_override.append(start_override)
        num.append(level_override)
        self.numbering.append(num)
        return num_id


def apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_element = OxmlElement("w:numId")
    num_id_element.set(qn("w:val"), str(num_id))
    num_pr.extend((ilvl, num_id_element))


def add_cover(doc: Document, metadata: GuideMetadata) -> None:
    section = doc.sections[0]
    configure_section(section, first=True)
    add_header_footer(section)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(58)

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(16)
    run = kicker.add_run("RELEASE OPERATIONS FIELD GUIDE")
    set_run_font(run, size=10.5)
    run.bold = True
    run.font.color.rgb = rgb(GOLD)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("혜니캘린더")
    set_run_font(run, size=30)
    run.font.color.rgb = rgb(NAVY)

    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(3)
    run = subtitle.add_run("Google Play 출시 가이드북")
    set_run_font(run, size=17)
    run.bold = True
    run.font.color.rgb = rgb(DEEP_BLUE)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_after = Pt(24)
    run = meta.add_run(
        f"버전 {metadata.version} · versionCode {metadata.version_code} · {metadata.written_date}"
    )
    set_run_font(run, size=10.5)
    run.font.color.rgb = rgb(GOLD)

    callout = doc.add_table(rows=1, cols=1)
    callout.style = "Table Grid"
    apply_table_geometry(callout, [CONTENT_WIDTH_DXA])
    cell = callout.cell(0, 0)
    shade(cell, PALE_BLUE)
    set_cell_margins(cell, 180, 260, 180, 260)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run("정책·보안 제출 차단 조건을 포함한 운영 기준")
    set_run_font(run, size=11)
    run.bold = True
    run.font.color.rgb = rgb(BLUE)
    detail = cell.add_paragraph()
    detail.alignment = WD_ALIGN_PARAGRAPH.CENTER
    detail.paragraph_format.space_after = Pt(0)
    run = detail.add_run("검증 증거가 없는 항목은 완료로 표시하지 않음")
    set_run_font(run, size=9.5)
    run.font.color.rgb = rgb(MUTED)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(58)
    prepared = doc.add_paragraph()
    prepared.alignment = WD_ALIGN_PARAGRAPH.CENTER
    prepared.paragraph_format.space_after = Pt(4)
    run = prepared.add_run(f"{metadata.package_name} 출시 운영용")
    set_run_font(run, size=11)
    run.bold = True
    run.font.color.rgb = rgb(NAVY)
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = note.add_run("민감정보·서명 비밀번호·실사용 가족 데이터는 포함하지 않음")
    set_run_font(run, size=9)
    run.italic = True
    run.font.color.rgb = rgb(MUTED)

    doc.add_page_break()


def should_keep_table_row_together(row_index: int, values: list[str]) -> bool:
    if row_index == 0:
        return True
    return sum(len(value.strip()) for value in values) <= TABLE_ROW_KEEP_TOGETHER_MAX_CHARS


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    apply_table_geometry(table, resolve_column_widths(column_count))

    for row_index, values in enumerate(rows):
        row_props = table.rows[row_index]._tr.get_or_add_trPr()
        if should_keep_table_row_together(row_index, values) and row_props.find(qn("w:cantSplit")) is None:
            row_props.append(OxmlElement("w:cantSplit"))
        for column_index in range(column_count):
            cell = table.cell(row_index, column_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index == 0:
                shade(cell, PALE_BLUE)
            value = values[column_index] if column_index < len(values) else ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.15
            add_inline(paragraph, value)
            for run in paragraph.runs:
                set_run_font(run, size=9)
                if row_index == 0:
                    run.bold = True
                    run.font.color.rgb = rgb(NAVY)

    header_props = table.rows[0]._tr.get_or_add_trPr()
    if header_props.find(qn("w:tblHeader")) is None:
        header_props.append(OxmlElement("w:tblHeader"))
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(1)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        parts = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell or "") for cell in parts):
            rows.append(parts)
        index += 1
    return rows, index


def add_body_from_markdown(doc: Document, text: str) -> None:
    lines = text.splitlines()
    numbering = NumberingManager(doc)
    active_list_kind: str | None = None
    active_num_id: int | None = None
    body_started = False
    index = 0

    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()

        if not body_started:
            if stripped.startswith("## "):
                body_started = True
            elif stripped and not is_cover_metadata_line(stripped):
                body_started = True
            else:
                index += 1
                continue

        if not stripped:
            active_list_kind = None
            active_num_id = None
            index += 1
            continue

        if stripped.startswith("|"):
            active_list_kind = None
            active_num_id = None
            rows, index = parse_table(lines, index)
            add_table(doc, rows)
            continue

        if stripped.startswith("```"):
            active_list_kind = None
            active_num_id = None
            index += 1
            code_lines: list[str] = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index].rstrip())
                index += 1
            code_paragraph = doc.add_paragraph(style="Code Block")
            shade(code_paragraph._p, "F5F7FA")
            run = code_paragraph.add_run("\n".join(code_lines))
            set_run_font(run, "Consolas", 8.5)
            index += 1
            continue

        heading = re.match(r"^(#{2,4})\s+(.+)$", stripped)
        if heading:
            active_list_kind = None
            active_num_id = None
            level = min(len(heading.group(1)) - 1, 3)
            paragraph = doc.add_paragraph(style=f"Heading {level}")
            add_inline(paragraph, heading.group(2))
            index += 1
            continue

        if stripped.startswith("> "):
            active_list_kind = None
            active_num_id = None
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.18)
            paragraph.paragraph_format.right_indent = Inches(0.18)
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(8)
            shade(paragraph._p, PALE_ROSE)
            add_inline(paragraph, stripped[2:])
            for run in paragraph.runs:
                set_run_font(run, size=10.5)
                run.bold = True
                run.font.color.rgb = rgb(ROSE)
            index += 1
            continue

        numbered = re.match(r"^(\d+)\.\s+(.+)$", stripped)
        bullet = re.match(r"^-\s+(.+)$", stripped)
        if numbered or bullet:
            kind = "decimal" if numbered else "bullet"
            if active_list_kind != kind or active_num_id is None:
                active_num_id = numbering.new_sequence(kind)
                active_list_kind = kind
            style = "List Number" if kind == "decimal" else "List Bullet"
            paragraph = doc.add_paragraph(style=style)
            apply_numbering(paragraph, active_num_id)
            add_inline(paragraph, numbered.group(2) if numbered else bullet.group(1))
            index += 1
            continue

        active_list_kind = None
        active_num_id = None
        paragraph = doc.add_paragraph()
        add_inline(paragraph, stripped)
        index += 1


def iter_table_paragraphs(table) -> Iterator:
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for nested in cell.tables:
                yield from iter_table_paragraphs(nested)


def iter_document_paragraphs(doc: Document) -> Iterator:
    yield from doc.paragraphs
    for table in doc.tables:
        yield from iter_table_paragraphs(table)
    for section in doc.sections:
        for container in (
            section.header,
            section.footer,
            section.even_page_header,
            section.even_page_footer,
        ):
            yield from container.paragraphs
            for table in container.tables:
                yield from iter_table_paragraphs(table)


def normalize_document(doc: Document) -> None:
    for paragraph in iter_document_paragraphs(doc):
        paragraph.paragraph_format.widow_control = True
        for run in paragraph.runs:
            if run.font.name is None:
                set_run_font(run)


def build(source: Path, output: Path) -> None:
    text = source.read_text(encoding="utf-8")
    metadata = parse_metadata(text)
    doc = Document()
    doc.core_properties.title = "혜니캘린더 Google Play 출시 가이드북"
    doc.core_properties.subject = "Google Play 출시 운영 절차와 정책 검증 체크리스트"
    doc.core_properties.author = "혜니캘린더"
    doc.core_properties.keywords = (
        f"Google Play, Android, 출시, 체크리스트, {PRESET_NAME}, {HEADER_TEMPLATE}"
    )
    doc.settings.odd_and_even_pages_header_footer = True
    configure_field_updates(doc)
    configure_styles(doc)
    add_cover(doc, metadata)
    add_body_from_markdown(doc, text)
    normalize_document(doc)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def parse_args(argv: list[str]) -> argparse.Namespace:
    default_source = find_default_source()
    parser = argparse.ArgumentParser(description="혜니캘린더 Google Play 출시 가이드 DOCX 생성")
    parser.add_argument(
        "source",
        nargs="?",
        type=Path,
        default=default_source,
        help=f"원본 Markdown (기본값: {default_source})",
    )
    parser.add_argument(
        "output",
        nargs="?",
        type=Path,
        help="출력 DOCX (기본값: 원본과 같은 이름의 .docx)",
    )
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    args = parse_args(sys.argv[1:] if argv is None else argv)
    source = args.source.resolve()
    output = args.output.resolve() if args.output else source.with_suffix(".docx")
    build(source, output)
    print(output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
